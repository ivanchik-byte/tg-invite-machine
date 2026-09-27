import io
from typing import Sequence
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from aiogram.types import BufferedInputFile

from app.models.models import Account

def generate_accounts_excel(accounts: Sequence[Account]) -> BufferedInputFile:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Аккаунты"

    headers = [
        "ID",
        "Телефон",
        "Имя",
        "Username",
        "Статус",
        "Активен",
        "Инвайтов сегодня",
        "Прокси",
        "Кулдаун до",
        "Последний инвайт"
    ]
    sheet.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = sheet.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    status_fills = {
        "active": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
        "spambot": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
        "cooldown": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
        "banned": PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid"),
    }

    for row_idx, acc in enumerate(accounts, start=2):
        proxy_display = "-"
        if acc.proxy:
            proxy_display = f"{acc.proxy.protocol}://{acc.proxy.host}:{acc.proxy.port}"

        last_invite_str = acc.last_invite_at.strftime("%Y-%m-%d %H:%M") if acc.last_invite_at else "-"
        cooldown_str = acc.cooldown_until.strftime("%Y-%m-%d %H:%M") if acc.cooldown_until else "-"

        row_data = [
            acc.id,
            acc.phone,
            acc.first_name or "-",
            f"@{acc.username}" if acc.username else "-",
            acc.status.upper(),
            "ДА" if acc.is_active else "НЕТ",
            acc.daily_invites_count,
            proxy_display,
            cooldown_str,
            last_invite_str,
        ]
        sheet.append(row_data)

        fill = status_fills.get(acc.status.lower())
        for col_idx in range(1, len(headers) + 1):
            cell = sheet.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if fill and col_idx == 5:
                cell.fill = fill
                cell.font = Font(bold=True)

    for col in sheet.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)

    return BufferedInputFile(
        file=output.read(),
        filename="accounts_audit.xlsx"
    )
